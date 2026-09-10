#!/usr/bin/env python3
"""
Regenerate digest/v2/data/backstory.json.

Source of truth:
  editorial/backstory-rows.json        — the 14 Lifetimes rows (hand-edited)
  editorial/NTK_Backstory_Object_Matrix.xlsx  — the object case
  FIRE (below)                         — the 7 Still Counting rows

Preserved from the existing backstory.json on every run:
  todays_pairings, and per-row: narrative, instances,
  updated, photo, indicator.as_of / .verified, and any hand-picked objects.

Run from the repo root:  python3 editorial/build_backstory.py
"""
import json, os, sys, collections
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROWS_IN = os.path.join(ROOT, 'editorial', 'backstory-rows.json')
MATRIX  = os.path.join(ROOT, 'editorial', 'NTK_Backstory_Object_Matrix.xlsx')
OUT     = os.path.join(ROOT, 'digest', 'v2', 'data', 'backstory.json')

# ── Still Counting ────────────────────────────────────────────────────
# Edit here. These seven have no start_line by design — the day count is
# the whole statement.
FIRE = [
 ('us-canada-trade','US–Canada trade war','2026-08-22',
  'Whether a trade fight between neighbours is leverage or self-harm.',
  'Both governments withdraw the tariffs imposed since August.'),
 ('us-israel-iran','US–Israel–Iran','2026-02-28',
  'Whether a campaign of strikes is a policy with an end, or a war nobody has named.',
  'A negotiated halt to strikes and retaliation lasting more than ninety days.'),
 ('venezuela','Venezuela, post-Maduro','2026-01-03',
  'Whether a state can be rebuilt by the people who fled it.',
  'A recognised government exercising control with international acceptance.'),
 ('israel-hezbollah','Israel–Hezbollah, Lebanon','2023-10-08',
  'Whether a border can be quiet without a settlement behind it.',
  'A durable withdrawal and ceasefire holding more than one year.'),
 ('gaza','Gaza','2023-10-07',
  'Whether there is a limit force cannot be argued past.',
  'A permanent ceasefire with an agreed administration for the territory.'),
 ('sudan','Sudan civil war','2023-04-15',
  'Whether a country can survive two armies that each believe they are the state.',
  'A signed settlement between the SAF and RSF holding six months.'),
 ('ukraine-russia','Ukraine–Russia','2022-02-24',
  'Whether borders can still be changed by force.',
  'A ceasefire or settlement recognised by both governments.'),
]

# ── Indicators ────────────────────────────────────────────────────────
# EVERY VALUE HERE IS UNVERIFIED. Check against the named source before
# publishing, then set "verified": true in backstory.json (the build
# preserves that flag).
IND = {
 'climate':       ('Atmospheric carbon dioxide', '351 ppm', '1988', '425 ppm', 'worse', 'NOAA Global Monitoring Laboratory'),
 'media':         ('Trust the press a great deal or fair amount', '72%', '1976', '31%', 'worse', 'Gallup'),
 'immigration':   ('Foreign-born share of the population', '4.7%', '1970', '15%', 'contested', 'U.S. Census Bureau'),
 'faith':         ('Adults with no religious affiliation', '7%', '1980', '28%', 'contested', 'Pew Research Center'),
 'taiwan':        ('PLA aircraft crossing the median line', 'about 0', '2019', 'over 3,000 a year', 'worse', 'Taiwan Ministry of National Defense'),
 'government':    ("Trust Washington to do what's right", '73%', '1958', '22%', 'worse', 'Pew Research Center'),
 'order':         ('Violent crime per 100,000 people', '758', '1991', '370', 'better', 'FBI Uniform Crime Reports'),
 'america-abroad':('U.S. troops stationed abroad', 'about 1 million', '1970', 'about 170,000', 'contested', 'Defense Manpower Data Center'),
 'power':         ('Say the other party threatens the nation', 'about 20%', '1994', 'about 70%', 'worse', 'Pew Research Center'),
 'work':          ('Share of workers in a union', '24%', '1973', '10%', 'worse', 'Bureau of Labor Statistics'),
 'family':        ('Median age at first marriage, women', '21', '1970', '28', 'contested', 'U.S. Census Bureau'),
 'equality':      ('Black family wealth per $100 of white family wealth', 'about $16', '1983', 'about $15', 'flat', 'Federal Reserve Survey of Consumer Finances'),
 'korea':         ('Estimated North Korean nuclear warheads', '0', '2005', 'about 50', 'worse', 'Federation of American Scientists'),
 'bomb':          ('Nuclear warheads worldwide', 'about 70,000', '1986', 'about 12,000', 'better', 'Federation of American Scientists'),
}

PRESERVE_ROW = ('narrative', 'instances', 'updated', 'photo')


def load_objects():
    sh = openpyxl.load_workbook(MATRIX)['Objects']
    by_row = collections.defaultdict(list)
    for i in range(2, sh.max_row + 1):
        g = lambda c: sh.cell(row=i, column=c).value
        if not g(11):
            continue
        by_row[g(11)].append({'title': g(6), 'author': g(7) or '',
                              'year': str(g(4))[:4] if g(4) else '',
                              'source': g(10), '_sort': str(g(4) or '9999')})
    return by_row


def pick(pool):
    """Two deepest pre-1981, two most recent. Falls back if the row is thin."""
    all_o = sorted(pool, key=lambda o: o['_sort'])
    if not all_o:
        return []
    old = [o for o in all_o if o['_sort'][:4] < '1981']
    new = [o for o in all_o if o['_sort'][:4] >= '1981']
    chosen = old[:2] + new[-2:]
    if len(chosen) < 4:
        chosen += [o for o in all_o if o not in chosen][:4 - len(chosen)]
    return [{k: o[k] for k in ('title', 'author', 'year', 'source')}
            for o in sorted(chosen, key=lambda o: o['_sort'])]


def main():
    spec = json.load(open(ROWS_IN))
    objs = load_objects()

    prev, prev_rows = {}, {}
    if os.path.exists(OUT):
        prev = json.load(open(OUT))
        prev_rows = {r['id']: r for r in prev.get('rows', [])}

    out = {
        '_comment': ('LIBRARY of 21 rows plus todays_pairings, the set actually rendered '
                     '(one entry per digest story). GENERATED by '
                     'editorial/build_backstory.py — edit editorial/backstory-rows.json '
                     'and rebuild. Narratives, instances, pairings and verified flags '
                     'ARE preserved across rebuilds.'),
        'version': '2.0',
        'generated': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'published': prev.get('published'),
        'rows': [],
        'todays_pairings': prev.get('todays_pairings', []),
    }

    def carry(row, old):
        for k in PRESERVE_ROW:
            if old and k in old:
                row[k] = old[k]
        # hand-picked objects win over the automatic pick
        if old and old.get('objects_locked'):
            row['objects'] = old['objects']
            row['objects_locked'] = True
        if row.get('indicator') and old and old.get('indicator'):
            for k in ('as_of', 'verified'):
                if k in old['indicator']:
                    row['indicator'][k] = old['indicator'][k]
        return row

    for rid, title, start, contest, close in FIRE:
        out['rows'].append(carry({
            'id': rid, 'stratum': 'fire', 'title': title,
            'start_date': start, 'start_line': None,
            'milestone': contest, 'close_condition': close,
            'roots': [], 'subgenres': [], 'updated': False,
            'narrative': [], 'indicator': None,
            'objects': [], 'instances': [], 'photo': None,
        }, prev_rows.get(rid)))

    for r in spec['rows']:
        ind = IND.get(r['id'])
        out['rows'].append(carry({
            'id': r['id'], 'stratum': 'held', 'title': r['title'],
            'start_date': r['start_date'], 'start_line': r['start_line'],
            'milestone': r['contest'], 'close_condition': None,
            'roots': r['roots'], 'subgenres': r['subgenres'],
            'updated': False, 'narrative': [],
            'indicator': ({'label': ind[0], 'then_value': ind[1], 'then_year': ind[2],
                           'now_value': ind[3], 'direction': ind[4], 'source': ind[5],
                           'as_of': None, 'verified': False} if ind else None),
            'objects': pick(objs.get(r['title'], [])),
            'instances': [], 'photo': None,
        }, prev_rows.get(r['id'])))

    out['rows'].sort(key=lambda r: r['start_date'], reverse=True)

    ids = [r['id'] for r in out['rows']]
    dupes = [i for i, c in collections.Counter(ids).items() if c > 1]
    if dupes:
        sys.exit('duplicate row ids: %s' % dupes)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w') as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
        f.write('\n')

    held = [r for r in out['rows'] if r['stratum'] == 'held']
    written = sum(1 for r in out['rows'] if r['narrative'])
    unver = sum(1 for r in held if r['indicator'] and not r['indicator']['verified'])
    print('wrote %s' % os.path.relpath(OUT, ROOT))
    print('  %d rows (%d fire, %d held)' % (len(out['rows']), len(FIRE), len(held)))
    print('  %d narratives written, %d empty' % (written, len(out['rows']) - written))
    print('  %d indicators UNVERIFIED' % unver)
    print('  %d pairings carried over' % len(out['todays_pairings']))


if __name__ == '__main__':
    main()
